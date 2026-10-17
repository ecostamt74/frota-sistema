import openpyxl
from datetime import datetime, date
from models import db, Motorista, Caminhao

def _to_date(valor):
    if not valor: return None
    if isinstance(valor, datetime): return valor.date()
    if isinstance(valor, date): return valor
    if isinstance(valor, str):
        valor = valor.split(' ')[0] # Pega só a data, ignora a hora
        for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
            try: return datetime.strptime(valor.strip(), fmt).date()
            except ValueError: continue
    return None

def _to_str(valor):
    return str(valor).strip() if valor is not None else ''

def importar_motoristas(caminho_xlsx):
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    if 'motoristas' not in wb.sheetnames: return "Aba 'motoristas' não encontrada."
    
    ws = wb['motoristas']
    cabecalho = [str(c.value).strip().lower() for c in ws[1]]
    importados = 0
    erros = 0
    
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if not any(linha): continue
        dados = dict(zip(cabecalho, linha))
        
        cpf = _to_str(dados.get('cpf'))
        if not cpf: continue
        
        # Se o CPF já existe, ele pula (para não dar erro de duplicidade)
        if Motorista.query.filter_by(cpf=cpf).first(): 
            continue
        
        try:
            m = Motorista(
                nome=_to_str(dados.get('nome')).upper(),
                cpf=cpf,
                cnh=_to_str(dados.get('cnh')),
                categoria=_to_str(dados.get('categoria')),
                validade_cnh=_to_date(dados.get('validade_cnh')),
                telefone=_to_str(dados.get('telefone')),
                origem='excel'
            )
            db.session.add(m)
            importados += 1
        except Exception as e:
            erros += 1
            print(f"Erro motorista: {e}")
        
    db.session.commit()
    return f"{importados} motorista(s) importado(s). {erros} erro(s)."

def importar_caminhoes(caminho_xlsx):
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    if 'caminhoes' not in wb.sheetnames: return "Aba 'caminhoes' não encontrada."
    
    ws = wb['caminhoes']
    cabecalho = [str(c.value).strip().lower() for c in ws[1]]
    importados = 0
    erros = 0
    
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if not any(linha): continue
        dados = dict(zip(cabecalho, linha))
        
        placa = _to_str(dados.get('placa')).upper().replace('-', '').replace(' ', '')
        if not placa: continue
        
        if Caminhao.query.filter_by(placa=placa).first(): continue
        
        try:
            c = Caminhao(
                placa=placa,
                marca=_to_str(dados.get('marca')),
                modelo=_to_str(dados.get('modelo')),
                ano=int(dados.get('ano')) if dados.get('ano') else None,
                km_atual=int(dados.get('km_atual') or 0),
                intervalo_troca_oleo=int(dados.get('intervalo_troca_oleo') or 10000),
                origem='excel'
            )
            db.session.add(c)
            importados += 1
        except Exception as e:
            erros += 1
            print(f"Erro caminhao: {e}")
            
    db.session.commit()
    return f"{importados} caminhão(ões) importado(s). {erros} erro(s)."